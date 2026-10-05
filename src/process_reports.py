import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

INPUT_DIR = Path("input")
OUTPUT_DIR = Path("output")
VALID_PROJECTS = {"P-101", "P-102", "P-103"}
MAX_DAILY_HOURS = 12

SYNONYMS = {
    "date": ["date", "work date", "day"],
    "project": ["project code", "project #", "project", "job code", "job #", "job"],
    "worker": ["worker", "employee", "name", "crew member"],
    "hours": ["hours", "hrs", "hours worked", "hrs worked"],
    "rate": ["rate", "hourly rate", "rate/hr", "rate per hour", "$/hr"],
    "description": ["description", "notes", "task", "work done"],
}
ISSUE_COLS = ["Subcontractor", "File", "Excel Row", "Problem", "Action"]
ISSUES = []


def add_issue(sub, file, row, problem, action):
    ISSUES.append(dict(zip(ISSUE_COLS, [sub, file, row, problem, action])))


def norm(v):
    return str(v).strip().lower() if pd.notna(v) else ""


def to_number(v):
    if pd.isna(v):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    m = re.search(r"\d+(\.\d+)?", str(v).replace(",", ""))
    return float(m.group()) if m else None


def clean_code(v):
    if pd.isna(v):
        return None
    s = re.sub(r"\s", "", str(v)).upper()
    m = re.fullmatch(r"P-?(\d{3})", s)
    return f"P-{m.group(1)}" if m else s


def parse_date(v):
    if pd.isna(v):
        return pd.NaT
    return pd.to_datetime(v, errors="coerce")


def load_file(path):
    raw = pd.read_excel(path, header=None, dtype=object)
    first = raw.iloc[0, 0] if not raw.empty else None
    sub = str(first).strip() if pd.notna(first) else path.stem

    header_idx, header_cells = None, None
    for idx in range(min(len(raw), 15)):
        cells = [norm(c) for c in raw.iloc[idx]]
        found = {f for f, names in SYNONYMS.items() if any(c in names for c in cells)}
        if len(found) >= 4:
            header_idx, header_cells = idx, cells
            break
    if header_idx is None:
        add_issue(sub, path.name, "", "Could not find a header row", "File skipped")
        return None

    rename = {}
    for field, names in SYNONYMS.items():
        for pos, c in enumerate(header_cells):
            if c in names:
                rename[pos] = field
                break

    body = raw.iloc[header_idx + 1:].copy()
    out = pd.DataFrame({field: body.iloc[:, pos] for pos, field in rename.items()})
    for field in SYNONYMS:
        if field not in out.columns:
            out[field] = None
    out["source_row"] = body.index + 1

    out = out.dropna(subset=list(SYNONYMS), how="all")
    if out.empty:
        return None
    is_total = out.drop(columns="source_row").apply(
        lambda r: any(str(v).strip().lower().startswith("total") for v in r if pd.notna(v)),
        axis=1,
    )
    out = out[~is_total].copy()
    out.insert(0, "sub", sub)
    out.insert(1, "file", path.name)
    return out


def style_workbook(path):
    wb = load_workbook(path)
    fill = PatternFill("solid", fgColor="1F3A5F")
    for ws in wb.worksheets:
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = fill
            cell.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"
        if ws.title == "Overview":
            for r in range(2, ws.max_row + 1):
                label = ws.cell(r, 1).value
                cell = ws.cell(r, 2)
                cell.alignment = Alignment(horizontal="left")
                if label == "Total cost":
                    cell.number_format = "$#,##0.00"
                elif label == "Total hours":
                    cell.number_format = "#,##0.0"
        for col in ws.columns:
            header = str(col[0].value)
            width = max(len(str(c.value)) if c.value is not None else 0 for c in col)
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max(width + 2, 10), 60)
            for c in col[1:]:
                if header in ("Cost", "Rate"):
                    c.number_format = "$#,##0.00"
                elif header == "Hours":
                    c.number_format = "0.0"
                elif header == "Date":
                    c.number_format = "yyyy-mm-dd"
    wb.save(path)

def add_formulas(path):
    wb = load_workbook(path)

    cd = wb["Clean Data"]  # A Sub, B Date, C Project, D Worker, E Hours, F Rate, G Cost
    for r in range(2, cd.max_row + 1):
        cd.cell(r, 7).value = f'=IF(OR(E{r}="",F{r}=""),"",E{r}*F{r})'

    ov = wb["Overview"]
    ov["B5"] = "=COUNTA('Clean Data'!A:A)-1"
    ov["B6"] = "=SUM('Clean Data'!E:E)"
    ov["B7"] = "=SUM('Clean Data'!G:G)"
    ov["B8"] = "=COUNTA(Exceptions!A:A)-1"

    bp = wb["By Project"]  # A Project, B Entries, C Hours, D Cost
    for r in range(2, bp.max_row + 1):
        bp.cell(r, 2).value = f"=COUNTIFS('Clean Data'!$C:$C,A{r})"
        bp.cell(r, 3).value = f"=SUMIFS('Clean Data'!$E:$E,'Clean Data'!$C:$C,A{r})"
        bp.cell(r, 4).value = f"=SUMIFS('Clean Data'!$G:$G,'Clean Data'!$C:$C,A{r})"

    bs = wb["By Subcontractor"]  # A Sub, B Entries, C Hours, D Cost, E Open Exceptions
    for r in range(2, bs.max_row + 1):
        bs.cell(r, 2).value = f"=COUNTIFS('Clean Data'!$A:$A,A{r})"
        bs.cell(r, 3).value = f"=SUMIFS('Clean Data'!$E:$E,'Clean Data'!$A:$A,A{r})"
        bs.cell(r, 4).value = f"=SUMIFS('Clean Data'!$G:$G,'Clean Data'!$A:$A,A{r})"
        bs.cell(r, 5).value = f"=COUNTIFS(Exceptions!$A:$A,A{r})"

    wb.save(path)


def main():
    files = sorted(p for p in INPUT_DIR.glob("*.xlsx") if not p.name.startswith("~$"))
    if not files:
        sys.exit("No .xlsx files in input/. Run generate_samples.py first.")

    frames = [f for f in (load_file(p) for p in files) if f is not None]
    df = pd.concat(frames, ignore_index=True)
    rows_read = len(df)

    df["date"] = pd.to_datetime(df["date"].apply(parse_date)).dt.normalize()
    df["project"] = df["project"].apply(clean_code)
    df["hours"] = df["hours"].apply(to_number)
    df["rate"] = df["rate"].apply(to_number)
    df["worker"] = df["worker"].apply(lambda v: str(v).strip() if pd.notna(v) else None)
    df["description"] = df["description"].apply(lambda v: str(v).strip() if pd.notna(v) else None)

    # Duplicates
    key = ["sub", "date", "project", "worker", "hours", "description"]
    dup = df.duplicated(subset=key, keep="first")
    for _, r in df[dup].iterrows():
        add_issue(r["sub"], r["file"], r["source_row"], "Duplicate entry", "Removed")
    df = df[~dup].copy()

    # Reporting week = the most common Monday-start week in the data
    week_start = (df["date"] - pd.to_timedelta(df["date"].dt.weekday, unit="D")).mode()[0]
    week_end = week_start + pd.Timedelta(days=6)

    excluded = pd.Series(False, index=df.index)
    for i, r in df.iterrows():
        s, f, row = r["sub"], r["file"], r["source_row"]
        if pd.isna(r["date"]):
            add_issue(s, f, row, "Missing or unreadable date", "Excluded")
            excluded[i] = True
        elif not (week_start <= r["date"] <= week_end):
            add_issue(s, f, row, f"Date {r['date']:%Y-%m-%d} is outside the reporting week", "Review")
        if pd.isna(r["hours"]):
            add_issue(s, f, row, "Missing hours", "Excluded")
            excluded[i] = True
        if not r["project"]:
            add_issue(s, f, row, "Missing project code", "Excluded")
            excluded[i] = True
        elif r["project"] not in VALID_PROJECTS:
            add_issue(s, f, row, f"Unknown project code '{r['project']}'", "Excluded")
            excluded[i] = True
        if pd.isna(r["rate"]):
            add_issue(s, f, row, "Missing rate", "Review")

    clean = df[~excluded].copy()

    daily = clean.groupby(["sub", "worker", "date"])["hours"].transform("sum")
    for i, r in clean[daily > MAX_DAILY_HOURS].iterrows():
        add_issue(r["sub"], r["file"], r["source_row"],
                  f"{r['worker']} logged {daily[i]:g} hours in one day", "Review")

    clean["cost"] = clean["hours"] * clean["rate"]
    exc = pd.DataFrame(ISSUES, columns=ISSUE_COLS)

    by_project = (clean.groupby("project")
                  .agg(Entries=("hours", "size"), Hours=("hours", "sum"), Cost=("cost", "sum"))
                  .reset_index().rename(columns={"project": "Project"}))
    by_sub = (clean.groupby("sub")
              .agg(Entries=("hours", "size"), Hours=("hours", "sum"), Cost=("cost", "sum"))
              .reset_index().rename(columns={"sub": "Subcontractor"}))
    counts = exc.groupby("Subcontractor").size()
    by_sub["Open Exceptions"] = by_sub["Subcontractor"].map(counts).fillna(0).astype(int)

    data = clean.rename(columns={
        "sub": "Subcontractor", "date": "Date", "project": "Project", "worker": "Worker",
        "hours": "Hours", "rate": "Rate", "cost": "Cost", "description": "Description",
        "file": "Source File", "source_row": "Source Row",
    })[["Subcontractor", "Date", "Project", "Worker", "Hours", "Rate", "Cost",
        "Description", "Source File", "Source Row"]].sort_values(["Date", "Subcontractor"])

    total_hours, total_cost = clean["hours"].sum(), clean["cost"].sum()
    overview = pd.DataFrame({
        "Item": ["Reporting week", "Files processed", "Rows read", "Rows in report",
                 "Total hours", "Total cost", "Exceptions flagged"],
        "Value": [f"{week_start:%Y-%m-%d} to {week_end:%Y-%m-%d}", len(frames), rows_read,
                  len(clean), round(float(total_hours), 1), round(float(total_cost), 2), len(exc)],
    })

    OUTPUT_DIR.mkdir(exist_ok=True)
    out_path = OUTPUT_DIR / f"weekly_project_report_{week_start:%Y-%m-%d}.xlsx"
    with pd.ExcelWriter(out_path, engine="openpyxl") as xw:
        overview.to_excel(xw, sheet_name="Overview", index=False)
        by_project.to_excel(xw, sheet_name="By Project", index=False)
        by_sub.to_excel(xw, sheet_name="By Subcontractor", index=False)
        data.to_excel(xw, sheet_name="Clean Data", index=False)
        exc.to_excel(xw, sheet_name="Exceptions", index=False)
    add_formulas(out_path)
    style_workbook(out_path)

    # Draft email (text only, nothing is sent)
    lines = [
        f"Subject: Weekly project hours and cost report, week of {week_start:%b %d, %Y}",
        "",
        "Hi team,",
        "",
        f"Attached is this week's combined report, built from {len(frames)} subcontractor timesheets.",
        "",
        f"- Total hours: {total_hours:,.1f}",
        f"- Total cost: ${total_cost:,.2f}",
        f"- Items flagged for review: {len(exc)}",
        "",
    ]
    if len(exc):
        lines.append("Top items to check:")
        for _, e in exc.head(5).iterrows():
            lines.append(f"- {e['Subcontractor']} (row {e['Excel Row']}): {e['Problem']} [{e['Action']}]")
        lines.append("")
    lines += ["Full details are on the Exceptions sheet.", "", "Thanks,", "Ashray"]
    (OUTPUT_DIR / "email_draft.txt").write_text("\n".join(lines), encoding="utf-8")

    print(f"Done. {len(frames)} files, {rows_read} rows read, {len(clean)} in report, {len(exc)} exceptions.")
    print(f"Report: {out_path}")


if __name__ == "__main__":
    main()